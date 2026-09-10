#!/usr/bin/env python3
"""Build a baseline-aware workbook for eval_mix_openml_aa_ter30_2607v1a."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


CONFIG_PATH = "recipe/phimm/config/v2607_new/eval_mix_openml_aa_ter30_2607v1a.yaml"
REFERENCE_MODEL = (
    "az://orngwus2cresco/data/boren/outputs/ver_2607/"
    "remax_2607v1_openml_verb_s100_bs256_lid/global_step_100/qwen_hf/"
)

GROUPS = {
    "mixlang": ["mixlang_fy26q2"],
    "openasr_ml": [
        "de_fleurs", "fr_fleurs", "it_fleurs", "es_fleurs", "pt_fleurs",
        "de_mcv", "fr_mcv", "it_mcv", "es_mcv", "fr_mls", "it_mls",
        "es_mls", "pt_mls",
    ],
    "AA": ["earnings22_cleaned_aa", "voxpopuli_cleaned_aa"],
    "inhouse": [
        "enus_conv_fy21q1", "enus_conv_om_fy25q3", "enus_dict_office_fy24q3",
        "dadk_conv_fy21q3", "dadk_conv_om_fy23q1", "dadk_dict_fy23q4",
        "huhu_conv_fy22q4", "huhu_conv_om_fy24q2", "huhu_dict_fy25q2",
        "nbno_conv_fy21q3", "nbno_conv_om_fy23q1", "nbno_dict_fy23q4",
        "nlnl_conv_fy23q2", "nlnl_conv_om_fy23q1", "nlnl_dict_fy23q4",
        "cscz_conv_fy23q2", "cscz_conv_om_fy24q2", "cscz_dict_fy24q2",
    ],
}

EXPECTED_METRIC = {
    dataset: ("dter_p_err" if group in {"mixlang", "inhouse"} else "p_err")
    for group, datasets in GROUPS.items()
    for dataset in datasets
}

METRIC_RE = re.compile(
    r"(?:val-core/)?(?P<dataset>[a-z0-9_]+)/"
    r"(?P<metric>dter_p_err|p_err)/mean@1['\"]?\s*(?:[:=]\s*|\s+)"
    r"(?P<value>[0-9.eE+\-]+)"
)


def _flatten_metrics(value: Any, prefix: str = "") -> dict[str, float]:
    flattened: dict[str, float] = {}
    if isinstance(value, dict):
        for key, child in value.items():
            child_prefix = f"{prefix}/{key}" if prefix else str(key)
            flattened.update(_flatten_metrics(child, child_prefix))
    elif isinstance(value, (int, float)):
        flattened[prefix] = float(value)
    return flattened


def load_metrics(source: Path) -> dict[str, float]:
    """Load final val-core metrics from a JSON object or captured Ray log."""
    text = source.read_text(encoding="utf-8")
    metrics: dict[str, float] = {}
    if source.suffix.lower() == ".json":
        flattened = _flatten_metrics(json.loads(text))
        for key, value in flattened.items():
            normalized = key.removeprefix("val-core/")
            parts = normalized.split("/")
            if len(parts) >= 3 and parts[-1] == "mean@1":
                dataset, metric = parts[-3], parts[-2]
                if EXPECTED_METRIC.get(dataset) == metric:
                    metrics[dataset] = value
    else:
        for match in METRIC_RE.finditer(text):
            dataset = match.group("dataset")
            metric = match.group("metric")
            if EXPECTED_METRIC.get(dataset) == metric:
                metrics[dataset] = float(match.group("value"))

    missing = sorted(set(EXPECTED_METRIC) - set(metrics))
    if missing:
        raise ValueError(f"{source} is missing {len(missing)} datasets: {', '.join(missing)}")
    return metrics


def build_workbook(
    baseline: dict[str, float],
    candidate: dict[str, float],
    baseline_label: str,
    candidate_label: str,
    candidate_model: str,
    baseline_source: str,
    candidate_source: str,
) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "results"
    sheet.append(["group", "dataset", "metric", baseline_label, candidate_label, "error reduction"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")

    group_rows: list[int] = []
    for group, datasets in GROUPS.items():
        first_row = sheet.max_row + 1
        for dataset in datasets:
            row = sheet.max_row + 1
            sheet.append([
                group,
                dataset,
                EXPECTED_METRIC[dataset],
                baseline[dataset],
                candidate[dataset],
                f"=1-E{row}/D{row}",
            ])
        last_row = sheet.max_row
        average_row = sheet.max_row + 1
        sheet.append([
            group,
            f"{group} avg",
            "arithmetic mean",
            f"=AVERAGE(D{first_row}:D{last_row})",
            f"=AVERAGE(E{first_row}:E{last_row})",
            f"=1-E{average_row}/D{average_row}",
        ])
        group_rows.append(average_row)
        for cell in sheet[average_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")

    overall_row = sheet.max_row + 1
    baseline_refs = ",".join(f"D{row}" for row in group_rows)
    candidate_refs = ",".join(f"E{row}" for row in group_rows)
    sheet.append([
        "overall",
        "overall avg",
        "mean of group averages",
        f"=AVERAGE({baseline_refs})",
        f"=AVERAGE({candidate_refs})",
        f"=1-E{overall_row}/D{overall_row}",
    ])
    for cell in sheet[overall_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")

    for row in range(2, sheet.max_row + 1):
        for column in range(4, 7):
            sheet.cell(row, column).number_format = "0.00%"
    sheet.conditional_formatting.add(
        f"F2:F{sheet.max_row}",
        ColorScaleRule(
            start_type="min", start_color="F8696B",
            mid_type="num", mid_value=0, mid_color="FFFFFF",
            end_type="max", end_color="63BE7B",
        ),
    )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:F{sheet.max_row}"
    for column, width in enumerate((22, 31, 20, 18, 18, 18), start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    provenance = workbook.create_sheet("provenance")
    provenance.append(["field", "value"])
    provenance.append(["config", CONFIG_PATH])
    provenance.append(["reference_model", REFERENCE_MODEL])
    provenance.append(["candidate_model", candidate_model])
    provenance.append(["baseline_metrics_source", baseline_source])
    provenance.append(["candidate_metrics_source", candidate_source])
    provenance.append(["metric_definition", "dter_p_err for MixLang/in-house; p_err for OpenASR-ML/AA"])
    provenance.append(["delta_definition", "1 - candidate_error / baseline_error"])
    provenance.column_dimensions["A"].width = 28
    provenance.column_dimensions["B"].width = 110
    return workbook


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--baseline", type=Path, required=True, help="Baseline metrics JSON or Ray log")
    parser.add_argument("--candidate", type=Path, required=True, help="Candidate metrics JSON or Ray log")
    parser.add_argument("--candidate-model", required=True)
    parser.add_argument("--candidate-label", default="candidate")
    parser.add_argument("--baseline-label", default="2607v1a reference")
    parser.add_argument("--out", type=Path, required=True)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    baseline = load_metrics(args.baseline)
    candidate = load_metrics(args.candidate)
    workbook = build_workbook(
        baseline,
        candidate,
        args.baseline_label,
        args.candidate_label,
        args.candidate_model,
        str(args.baseline),
        str(args.candidate),
    )
    args.out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.out)
    print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()