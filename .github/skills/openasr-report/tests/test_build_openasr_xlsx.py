import importlib.util
import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook


SCRIPT = Path(__file__).parents[1] / "scripts/build_openasr_xlsx.py"
SUPPLIED = {
    "ami_clean": (6.72, 6.78),
    "earnings22-cleaned-aa-chunked": (5.02, 5.37),
    "gigaspeech_clean": (7.58, 7.52),
    "librispeech-clean": (1.26, 1.27),
    "librispeech-other": (2.74, 2.88),
    "monsoon_en_in": (3.39, 3.41),
    "spgispeech": (2.15, 2.13),
    "voxpopuli-cleaned-aa": (1.63, 1.69),
    "urgent2024": (9.57, 9.00),
    "urgent2024_clean": (4.51, 4.00),
    "de_fleurs": (2.32, 2.24), "de_mcv": (1.90, 1.90),
    "es_fleurs": (2.71, 2.81), "es_mcv": (2.20, 2.16), "es_mls": (2.58, 2.67),
    "fr_fleurs": (2.86, 2.65), "fr_mcv": (4.17, 4.18), "fr_mls": (2.37, 2.43),
    "monsoon_hi_in": (8.78, 9.63),
    "it_fleurs": (1.30, 1.27), "it_mcv": (1.78, 1.85), "it_mls": (3.99, 4.10),
    "nl_fleurs": (3.88, 3.65), "nl_mcv": (1.64, 1.81), "nl_mls": (4.53, 3.99),
    "pt_fleurs": (3.10, 3.41), "pt_mls": (3.36, 3.27),
    "earnings21": (7.9156, 7.0),
    "earnings22": (10.1216, 11.0),
    "coraal": (15.4380, 14.0),
    "earnings22_cleaned_aa": (4.3131, 4.0),
}


@pytest.fixture
def report():
    spec = importlib.util.spec_from_file_location("build_openasr_xlsx", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_baseline_and_full_report(report, tmp_path, monkeypatch):
    assert report.BASELINE_LABEL == "2609r2"
    baseline = {key: values[0] / 100 for key, values in SUPPLIED.items()}
    candidate = {key: values[1] / 100 for key, values in SUPPLIED.items()}
    assert report.BASELINE_METRICS == pytest.approx(baseline)
    assert set(report.OPENASR_DATASETS) | {
        key for _, datasets in report.OPENASR_ML_GROUPS for key in datasets
    } | set(report.OPENLONG_DATASETS) == set(SUPPLIED)
    source = tmp_path / "candidate.json"
    source.write_text(json.dumps(candidate))
    output = tmp_path / "report.xlsx"
    monkeypatch.setattr(
        sys, "argv", [str(SCRIPT), "B", "--metrics", str(source), "--out", str(output)],
    )
    assert report.main() == 0
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["openasr", "model_info", "dataset_results"]
    sheet = workbook["openasr"]
    assert sheet["A3"].value == "OpenEN"
    assert sheet["A15"].value == "OpenML"
    seen = set()
    for row in sheet.iter_rows():
        name = row[0].value
        if name in SUPPLIED:
            seen.add(name)
            assert row[1].value == pytest.approx(baseline[name])
            assert row[2].value == pytest.approx(candidate[name])
            assert row[3].value == f"=1-C{row[0].row}/B{row[0].row}"
            assert all(cell.number_format == "0.00%" for cell in row[1:4])
    assert seen == set(SUPPLIED)
    assert sheet["B14"].value == pytest.approx(44.57 / 1000)
    assert sheet["C14"].value == pytest.approx(44.05 / 1000)
    assert sheet["B40"].value == pytest.approx(3.6366666666666667 / 100)
    candidate_ml_mean = sum(
        sum(candidate[ds] for ds in datasets) / len(datasets)
        for _, datasets in report.OPENASR_ML_GROUPS
    ) / len(report.OPENASR_ML_GROUPS)
    assert sheet["C40"].value == pytest.approx(candidate_ml_mean)
    assert sheet["D40"].value == pytest.approx(1 - candidate_ml_mean / sheet["B40"].value)
    assert sheet["B46"].value == pytest.approx(0.09447075)
    assert sheet["C46"].value == pytest.approx(0.09)
    assert sheet["D46"].value == pytest.approx(1 - 0.09 / 0.09447075)
    summaries = {
        "avg": "4.46%", "de avg": "2.11%", "es avg": "2.50%",
        "fr avg": "3.13%", "hi avg": "8.78%", "it avg": "2.36%",
        "nl avg": "3.35%", "pt avg": "3.23%", "ml avg": "3.64%",
    }
    for row in sheet:
        if row[0].value in summaries:
            assert f"{row[1].value:.2%}" == summaries[row[0].value]
    workbook.close()
    columns = report.read_existing_xlsx(output)
    assert columns[0][0] == report.BASELINE_LABEL
    assert columns[0][1] == pytest.approx(baseline)
    assert columns[1][1] == pytest.approx(candidate)


def test_log_names_and_latest_value(report):
    text = "\n".join([
        "val-aux/earnings22-cleaned-aa-chunked/p_err/mean@1:0.0537",
        "val-aux/fleurs_nl/p_err/mean@1:0.04",
        "val-aux/nl_fleurs/p_err/mean@1:0.0365",
        "val-aux/monsoon_hi_in/p_err/mean@1:0.0963",
        "val-aux/urgent2024/p_err/mean@1:0.09",
        "val-aux/urgent2024_clean/p_err/mean@1:0.04",
    ])
    assert report.parse_wer_lines(text) == {
        "earnings22-cleaned-aa-chunked": 0.0537,
        "nl_fleurs": 0.0365,
        "monsoon_hi_in": 0.0963,
        "urgent2024": 0.09,
        "urgent2024_clean": 0.04,
    }


def test_json_aliases_and_baseline_override(report, tmp_path, monkeypatch):
    source = tmp_path / "metrics.json"
    source.write_text(json.dumps({"fleurs_nl": 0.10, "mcv_nl": 0.20, "mls_nl": 0.30}))
    output = tmp_path / "override.xlsx"
    monkeypatch.setattr(
        sys, "argv",
        [str(SCRIPT), "B", "--metrics", str(source), "--baseline", str(source),
         "--baseline-label", "explicit", "--out", str(output)],
    )
    assert report.main() == 0
    columns = report.read_existing_xlsx(output)
    assert columns[0][0] == "explicit"
    assert columns[0][1] == columns[1][1] == {"nl_fleurs": 0.10, "nl_mcv": 0.20, "nl_mls": 0.30}


def test_rejects_old_dataset_layout(report, tmp_path):
    output = tmp_path / "old.xlsx"
    report.build_workbook([(report.BASELINE_LABEL, report.BASELINE_METRICS)], output)
    workbook = load_workbook(output)
    workbook["openasr"]["A4"] = "ami"
    workbook.save(output)
    workbook.close()
    with pytest.raises(ValueError, match="incompatible dataset layout at row 4"):
        report.read_existing_xlsx(output)


def test_partial_coverage_deltas_use_matched_datasets(report, tmp_path):
    baseline = {"ami_clean": 0.10, "spgispeech": 0.40, "de_fleurs": 0.10, "de_mcv": 0.30}
    candidate = {"ami_clean": 0.05, "de_fleurs": 0.20}
    output = tmp_path / "partial.xlsx"
    report.build_workbook([("baseline", baseline), ("candidate", candidate)], output)
    workbook = load_workbook(output)
    sheet = workbook["openasr"]
    assert sheet["D4"].value == "=1-C4/B4"
    assert sheet["D14"].value == pytest.approx(0.50)
    assert sheet["B14"].value == pytest.approx(0.25)
    assert sheet["C14"].value == pytest.approx(0.05)
    assert sheet["D14"].comment is None
    assert sheet["D16"].value == "=1-C16/B16"
    assert sheet["D18"].value == pytest.approx(-1.0)
    assert sheet["D40"].value == pytest.approx(-1.0)
    assert sheet["D40"].comment is None
    for row in (5, 10, 12, 13, 17):
        assert sheet.cell(row, 4).value is None
        assert sheet.cell(row, 4).comment is None
    assert sheet.max_row == 46
    assert all(cell.comment is None for row in sheet for cell in row)
    assert all(cell.value is None for cell in sheet[1])
    workbook.close()


def test_zero_baseline_and_zero_candidate(report, tmp_path):
    output = tmp_path / "zeros.xlsx"
    report.build_workbook([
        ("baseline", {"ami_clean": 0.0, "spgispeech": 0.10, "de_fleurs": 0.0}),
        ("candidate", {"ami_clean": 0.05, "spgispeech": 0.0, "de_fleurs": 0.0}),
    ], output)
    workbook = load_workbook(output)
    sheet = workbook["openasr"]
    assert sheet["D4"].value is None
    assert sheet["D4"].comment is None
    assert sheet["D10"].value == "=1-C10/B10"
    assert sheet["C10"].value == 0
    assert sheet["D14"].value == pytest.approx(0.5)
    assert sheet["D40"].value is None
    assert sheet["D40"].comment is None
    workbook.close()


def test_ml_macro_mean_and_werr_use_matched_language_coverage(report, tmp_path):
    baseline = {
        "de_fleurs": 0.10, "de_mcv": 0.30,
        "fr_fleurs": 0.60, "fr_mcv": 0.80, "fr_mls": 0.90,
        "monsoon_hi_in": 0.90,
    }
    candidate = {"de_fleurs": 0.05, "fr_fleurs": 0.15, "fr_mcv": 0.20}
    output = tmp_path / "macro.xlsx"
    report.build_workbook([("baseline", baseline), ("candidate", candidate)], output)
    workbook = load_workbook(output)
    sheet = workbook["openasr"]
    assert sheet["B40"].value == pytest.approx((0.20 + 2.30 / 3 + 0.90) / 3)
    assert sheet["C40"].value == pytest.approx((0.05 + 0.175) / 2)
    assert sheet["D40"].value == pytest.approx(1 - 0.1125 / 0.40)
    assert sheet["D40"].value != pytest.approx(1 - 0.40 / 1.50)
    workbook.close()


def test_no_common_ml_language_leaves_werr_blank(report, tmp_path):
    output = tmp_path / "unmatched.xlsx"
    report.build_workbook([
        ("baseline", {"de_fleurs": 0.10}),
        ("candidate", {"fr_fleurs": 0.20}),
    ], output)
    workbook = load_workbook(output)
    sheet = workbook["openasr"]
    assert sheet["B40"].value == pytest.approx(0.10)
    assert sheet["C40"].value == pytest.approx(0.20)
    assert sheet["D40"].value is None
    workbook.close()


def test_every_model_gets_its_own_delta(report, tmp_path, monkeypatch):
    prior = tmp_path / "prior.xlsx"
    baseline = {"ami_clean": 0.10}
    report.build_workbook([("baseline", baseline), ("better", {"ami_clean": 0.05})], prior)
    candidate = tmp_path / "candidate.json"
    candidate.write_text(json.dumps({"ami_clean": 0.20}))
    output = tmp_path / "extended.xlsx"
    monkeypatch.setattr(sys, "argv", [
        str(SCRIPT), "worse", "--metrics", str(candidate),
        "--extend-xlsx", str(prior), "--out", str(output),
    ])
    assert report.main() == 0
    workbook = load_workbook(output)
    sheet = workbook["openasr"]
    assert sheet["E4"].value == "=1-C4/B4"
    assert sheet["F4"].value == "=1-D4/B4"
    assert sheet["E14"].value == pytest.approx(0.5)
    assert sheet["F14"].value == pytest.approx(-1.0)
    assert len(sheet.conditional_formatting) == 2
    for key in sheet.conditional_formatting:
        rules = sheet.conditional_formatting[key]
        assert rules[0].colorScale.cfvo[1].val == 0
    workbook.close()


@pytest.mark.parametrize("value", [float("nan"), float("inf"), -0.1])
def test_invalid_error_rate_is_rejected(report, tmp_path, value):
    with pytest.raises(ValueError, match="finite, nonnegative"):
        report.build_workbook([("baseline", {"ami_clean": value})], tmp_path / "invalid.xlsx")


def test_metadata_round_trip_and_extension(report, tmp_path, monkeypatch):
    metrics = tmp_path / "metrics.json"
    metrics.write_text(json.dumps({
        "ami_clean": 0.05, "de_fleurs": 0.02,
        "urgent2024": 0.09, "urgent2024_clean": 0.04,
    }))
    model_path = tmp_path / "model.json"
    result_path = tmp_path / "results.json"
    prior = tmp_path / "prior.xlsx"
    first = {
        "model_path": "az://account/container/model1/qwen_hf",
        "config": "recipe/phimm/config/eval/example.yaml",
        "node": "verl-n4-i4",
        "ray_job_id": "raysubmit_first",
        "wandb_url": "https://wandb.example/project/runs/first",
        "code_snapshot": "abc123",
    }
    first_results = {
        "ami_clean": "az://account/container/model1/ami/0.jsonl",
        "de_fleurs": "/results/model1/de_fleurs/result_details.jsonl",
        "urgent2024": "/results/model1/urgent2024/0.jsonl",
        "urgent2024_clean": "/results/model1/urgent2024_clean/0.jsonl",
    }
    model_path.write_text(json.dumps(first))
    result_path.write_text(json.dumps({**first_results, "fleurs_de": first_results["de_fleurs"]}))
    monkeypatch.setattr(sys, "argv", [
        str(SCRIPT), "first", "--metrics", str(metrics),
        "--model-info", str(model_path), "--dataset-results", str(result_path), "--out", str(prior),
    ])
    assert report.main() == 0
    info, results = report.read_existing_metadata(prior)
    assert info["first"] == first
    assert info["2609r2"] == {}
    assert results == {"first": first_results}
    second = {"model_path": "/models/second", "ray_job_id": "raysubmit_second"}
    second_results = {"ami_clean": "/results/second/ami/0.jsonl"}
    model_path.write_text(json.dumps(second))
    result_path.write_text(json.dumps(second_results))
    output = tmp_path / "extended.xlsx"
    monkeypatch.setattr(sys, "argv", [
        str(SCRIPT), "second", "--metrics", str(metrics), "--extend-xlsx", str(prior),
        "--model-info", str(model_path), "--dataset-results", str(result_path), "--out", str(output),
    ])
    assert report.main() == 0
    info, results = report.read_existing_metadata(output)
    assert info == {"2609r2": {}, "first": first, "second": second}
    assert results == {"first": first_results, "second": second_results}
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["openasr", "model_info", "dataset_results"]
    assert workbook["openasr"].max_row == 46
    assert workbook["model_info"].max_row == 4
    assert workbook["dataset_results"].max_row == 94
    assert all(cell.comment is None for sheet in workbook for row in sheet for cell in row)
    workbook.close()


def test_metric_only_workbook_metadata_is_optional(report, tmp_path):
    path = tmp_path / "legacy.xlsx"
    report.build_workbook([("baseline", {"ami_clean": 0.1})], path)
    workbook = load_workbook(path)
    del workbook["model_info"]
    del workbook["dataset_results"]
    workbook.save(path)
    workbook.close()
    assert report.read_existing_metadata(path) == ({}, {})
    assert report.read_existing_xlsx(path) == [("baseline", {"ami_clean": 0.1})]


def test_unknown_metadata_is_rejected(report, tmp_path):
    with pytest.raises(ValueError, match="unknown model information fields"):
        report.build_workbook([("candidate", {})], tmp_path / "bad.xlsx",
                              model_info={"candidate": {"model_paht": "wrong key"}})
    with pytest.raises(ValueError, match="unknown dataset result names"):
        report.build_workbook([("candidate", {})], tmp_path / "bad.xlsx",
                              dataset_results={"candidate": {"wrong_dataset": "/some/result.jsonl"}})


def test_openlong_logs_metadata_and_extension(report, tmp_path, monkeypatch):
    baseline = {ds: report.BASELINE_METRICS[ds] for ds in report.OPENLONG_DATASETS}
    candidate = {ds: value / 2 for ds, value in baseline.items()}
    logs = tmp_path / "eval.log"
    logs.write_text("\n".join(
        f"val-aux/{ds}/p_err/mean@1:{value}" for ds, value in candidate.items()
    ))
    results = {
        ds: f"az://account/container/candidate/val_data_gen/{ds}/0.jsonl"
        for ds in candidate
    }
    result_file = tmp_path / "results.json"
    result_file.write_text(json.dumps(results))
    prior = tmp_path / "openlong.xlsx"
    monkeypatch.setattr(sys, "argv", [
        str(SCRIPT), "candidate", "--from-text", str(logs),
        "--dataset-results", str(result_file), "--out", str(prior),
    ])
    assert report.main() == 0
    assert report.read_existing_metadata(prior)[1] == {"candidate": results}
    assert report.read_existing_xlsx(prior)[1] == ("candidate", candidate)
    output = tmp_path / "extended.xlsx"
    monkeypatch.setattr(sys, "argv", [
        str(SCRIPT), "next", "--from-text", str(logs),
        "--extend-xlsx", str(prior), "--out", str(output),
    ])
    assert report.main() == 0
    assert report.read_existing_metadata(output)[1] == {"candidate": results}
    workbook = load_workbook(output)
    sheet = workbook["openasr"]
    assert sheet["A41"].value == "OpenLong"
    assert sheet["A46"].value == "openlong avg"
    assert sheet["B46"].value == pytest.approx(sum(baseline.values()) / 4)
    for row in range(42, 46):
        assert sheet.cell(row, 5).value == f"=1-C{row}/B{row}"
        assert sheet.cell(row, 6).value == f"=1-D{row}/B{row}"
    for column in range(2, 7):
        cell = sheet.cell(46, column)
        assert cell.number_format == "0.00%"
        assert cell.alignment.horizontal == "center"
        assert cell.fill == report.OVERALL_AVG_FILL
    assert sheet["E46"].value == pytest.approx(0.5)
    assert sheet["F46"].value == pytest.approx(0.5)
    assert sheet["C14"].value is None
    assert sheet["C40"].value is None
    for key in sheet.conditional_formatting:
        assert str(key.sqref).endswith("46")
        assert sheet.conditional_formatting[key][0].colorScale.cfvo[1].val == 0
    workbook.close()


@pytest.mark.parametrize("candidate, expected_werr", [
    ({"earnings21": 0.05}, 0.5),
    ({"earnings21": 0.20}, -1.0),
    ({"earnings21": 0.0}, 1.0),
    ({"coraal": 0.05}, None),
])
def test_openlong_partial_coverage_and_zero_baseline(report, tmp_path, candidate, expected_werr):
    output = tmp_path / "partial.xlsx"
    report.build_workbook([
        ("baseline", {"earnings21": 0.10, "earnings22": 0.30, "coraal": 0.0}),
        ("candidate", candidate),
    ], output)
    workbook = load_workbook(output)
    sheet = workbook["openasr"]
    assert sheet["B46"].value == pytest.approx(0.40 / 3)
    assert sheet["C46"].value == pytest.approx(next(iter(candidate.values())))
    if expected_werr is None:
        assert sheet["D46"].value is None
    else:
        assert sheet["D46"].value == pytest.approx(expected_werr)
        assert sheet["D42"].value == "=1-C42/B42"
    assert sheet["D43"].value is None
    assert sheet["D44"].value is None
    assert sheet["D45"].value is None
    workbook.close()


def test_extend_pre_openlong_report_preserves_existing_values(report, tmp_path, monkeypatch):
    prior = tmp_path / "legacy.xlsx"
    columns = [("baseline", {"ami_clean": 0.1}), ("first", {"ami_clean": 0.05})]
    report.build_workbook(columns, prior)
    workbook = load_workbook(prior)
    workbook["openasr"].delete_rows(41, 6)
    workbook.save(prior)
    workbook.close()
    assert report.read_existing_xlsx(prior) == columns
    metrics = tmp_path / "metrics.json"
    metrics.write_text(json.dumps({"earnings21": 0.05}))
    output = tmp_path / "extended.xlsx"
    monkeypatch.setattr(sys, "argv", [
        str(SCRIPT), "long", "--metrics", str(metrics),
        "--extend-xlsx", str(prior), "--out", str(output),
    ])
    assert report.main() == 0
    assert report.read_existing_xlsx(output) == columns + [("long", {"earnings21": 0.05})]
    workbook = load_workbook(output)
    assert workbook["openasr"]["B42"].value is None
    assert workbook["openasr"]["F42"].value is None
    workbook.close()


def test_rejects_mismatched_openlong_layout(report, tmp_path):
    output = tmp_path / "bad.xlsx"
    report.build_workbook([("baseline", report.BASELINE_METRICS)], output)
    workbook = load_workbook(output)
    workbook["openasr"]["A45"] = "earnings22-cleaned-aa-chunked"
    workbook.save(output)
    workbook.close()
    with pytest.raises(ValueError, match="incompatible dataset layout at row 45"):
        report.read_existing_xlsx(output)
